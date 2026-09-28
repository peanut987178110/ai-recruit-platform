"""简历解析服务（R-02、R-08）。

三层注入防护（PRD 4.5）在这里落地前两层：
  输入层 —— 指令特征剥离：识别并剥离疑似指令文本，原文留档标记待审
  提示层 —— 内容边界隔离：在提示词里声明简历是数据不是指令
输出层的一致性校验在打分服务里做。
"""
from __future__ import annotations

import re
import time
from pathlib import Path

from app.core.schemas import ResumeParseResult
from app.llm.client import llm
from app.prompts.manager import prompts

# 指令特征：中英文的「忽略指令」「给我满分」「角色扮演」等模式
_INJECTION_PATTERNS = [
    re.compile(r"忽略(上述|以上|前面|之前)?.{0,6}(要求|指令|规则|提示)"),
    re.compile(r"(ignore|disregard)\s+(all\s+)?(previous|above|prior)\s+(instructions?|prompts?|rules?)", re.I),
    re.compile(r"(给|打)(我|他|该候选人)?\s*(满分|最高分|100\s*分)"),
    re.compile(r"(score|rate|give)\s+(me|him|her|this\s+candidate)?\s*(full|max|100)\s*(score|marks?|points?)", re.I),
    re.compile(r"你(现在)?(是|扮演|充当).{0,12}(助手|专家|面试官|AI)"),
    re.compile(r"(you\s+are|act\s+as|pretend\s+to\s+be).{0,20}(assistant|expert|interviewer|ai)", re.I),
    re.compile(r"(输出|返回|打印).{0,10}(时)?.{0,6}(不|无需|不要).{0,8}(校验|检查|验证)"),
    re.compile(r"(system\s*prompt|系统提示词|提示词注入|prompt\s*injection)", re.I),
    re.compile(r"(务必|必须|一定).{0,6}(通过|录取|录用)(该|此|这名)?候选人"),
    re.compile(r"<\s*(system|instruction|prompt)\s*>", re.I),
]

# 疑似隐藏文本：零宽字符、白字等常见注入载体
_HIDDEN_CHARS = re.compile(r"[​-‏‪-‮⁠-⁯﻿]")


def detect_injection(text: str) -> list[str]:
    """返回命中的疑似指令片段原文。命中不拒绝，只剥离并标记待审（PRD 4.5）。"""
    hits: list[str] = []
    for pat in _INJECTION_PATTERNS:
        for m in pat.finditer(text):
            snippet = text[max(0, m.start() - 20): m.end() + 20].replace("\n", " ").strip()
            hits.append(snippet)
    if _HIDDEN_CHARS.search(text):
        hits.append("检测到零宽字符或双向控制字符（常见于隐藏文字注入）")
    return hits[:8]


def strip_injection(text: str) -> tuple[str, list[str]]:
    """剥离疑似指令文本，返回 (清理后文本, 命中列表)。原文由调用方留档。"""
    hits = detect_injection(text)
    cleaned = text
    for pat in _INJECTION_PATTERNS:
        cleaned = pat.sub("［已剥离的疑似指令文本］", cleaned)
    cleaned = _HIDDEN_CHARS.sub("", cleaned)
    return cleaned, hits


# ---------------- 文件文本抽取 ----------------


def extract_text(path: Path, file_type: str) -> str:
    """从 PDF / Word / 图片 / 文本中抽取纯文本。

    图片走视觉模型（见 parse_image），这里只处理有文本层的格式。
    """
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n".join((p.extract_text() or "") for p in reader.pages).strip()
        if suffix in (".docx", ".doc"):
            import docx
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts).strip()
        if suffix in (".txt", ".md", ".json", ".csv"):
            for enc in ("utf-8", "gbk", "utf-16"):
                try:
                    return path.read_text(encoding=enc).strip()
                except UnicodeDecodeError:
                    continue
            return path.read_bytes().decode("utf-8", errors="ignore").strip()
        if suffix in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines).strip()
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"文件无法打开，请重新上传（{type(e).__name__}）") from e
    raise ValueError(f"不支持的文件类型：{suffix}")


def _mask_school(school: str) -> str:
    """院校名称脱敏后单独存储，不参与打分（PRD 3.2.2）。"""
    if not school:
        return ""
    s = school.strip()
    if len(s) <= 2:
        return "＊" * len(s)
    return s[0] + "＊" * (len(s) - 2) + s[-1]


def _keyword_fallback(text: str) -> ResumeParseResult:
    """模型不可用时的关键词兜底解析。

    这不是静默降级：返回结果里 model 标记为空、降级原因会写进界面。
    兜底只保证字段有基本可用值，准确率远低于模型路径，界面必须提示人工核对。
    """
    res = ResumeParseResult(raw_text=text)
    if m := re.search(r"(?:姓名|Name)[:：\s]*([一-鿿]{2,4})", text):
        res.name = m.group(1)
    elif m := re.match(r"^\s*([一-鿿]{2,4})\s*$", text.split("\n")[0] if text else ""):
        res.name = m.group(1)
    if m := re.search(r"1[3-9]\d{9}", text):
        res.phone = m.group(0)
    if m := re.search(r"[\w.\-+]+@[\w\-]+\.[\w.]+", text):
        res.email = m.group(0)
    for lvl in ("博士", "硕士", "本科", "大专", "专科", "高中"):
        if lvl in text:
            res.education_level = lvl
            break
    if m := re.search(r"(?:院校|学校|大学|学院)[:：\s]*([一-鿿]{2,12}(?:大学|学院|学校))", text):
        res.school_masked = _mask_school(m.group(1))
    skills = ["Java", "Go", "Python", "SQL", "Redis", "Kafka", "MySQL", "Vue", "React",
              "TypeScript", "Excel", "Tableau", "Spring Boot", "Kubernetes"]
    res.skills = [s for s in skills if s.lower() in text.lower()]
    res.metrics = re.findall(r"[^\n。；]{0,20}\d+(?:\.\d+)?\s*(?:%|倍|万|QPS|ms|秒|小时|天)[^\n。；]{0,10}", text)[:8]
    return res


async def parse_resume(
    text: str,
    candidate_id: int | None = None,
    images: list[str] | None = None,
) -> tuple[ResumeParseResult, dict]:
    """解析简历。返回 (结构化结果, 元信息字典)。"""
    started = time.time()

    # 输入层防护：剥离疑似指令，原文留档
    original = text
    cleaned, hits = strip_injection(text)
    meta: dict = {
        "injection_hits": hits,
        "original_kept": bool(hits),
        "degraded": False,
        "path": "model",
    }

    p = await prompts.resolve("简历解析", {"resume_text": cleaned[:12000]},
                              route_key=str(candidate_id or cleaned[:60]))

    schema = ResumeParseResult
    res = await llm.complete_json(
        ability="简历解析", prompt_id=p.prompt_id, prompt_version=p.version,
        system=p.system_prompt, user=p.user_prompt, schema=schema,
        tier=p.tier, max_tokens=4000, is_async=True,
        images=images,
    )

    if res.ok and res.result:
        out: ResumeParseResult = res.result
        out.raw_text = original
        out.injection_hits = hits
        # 院校名称脱敏（模型可能返回 school_masked 或原始校名）
        if out.school_masked and "＊" not in out.school_masked:
            out.school_masked = _mask_school(out.school_masked)
        meta.update({
            "model": res.meta.model, "latency_ms": res.meta.latency_ms,
            "degrade": res.meta.degrade.value, "confidence": res.meta.confidence,
            "cost_cny": res.meta.cost_cny, "prompt_version": res.meta.prompt_version,
            "tokens_in": res.meta.tokens_in, "tokens_out": res.meta.tokens_out,
        })
        return out, meta

    # 降级：关键词兜底，明确标记
    fb = _keyword_fallback(cleaned)
    fb.raw_text = original
    fb.injection_hits = hits
    meta.update({
        "degraded": True, "path": "fallback",
        "degrade_reason": res.error or "模型调用失败",
        "degrade": res.meta.degrade.value,
        "latency_ms": int((time.time() - started) * 1000),
    })
    return fb, meta


def check_missing_fields(parsed: ResumeParseResult) -> list[str]:
    """关键字段缺失检查（PRD 3.2.2 兜底矩阵）。返回缺失字段的中文名。"""
    missing = []
    if not parsed.name:
        missing.append("姓名")
    if not parsed.phone and not parsed.email:
        missing.append("联系方式")
    if not parsed.works:
        missing.append("工作经历")
    if not parsed.education_level:
        missing.append("教育背景")
    return missing


def spot_multiple_resumes(text: str) -> int:
    """检测单文件多人简历（PRD 3.2.2）。返回疑似简历份数。"""
    markers = len(re.findall(r"(?:求职意向|应聘职位|期望职位|求职岗位)[:：]", text))
    phone_count = len(set(re.findall(r"1[3-9]\d{9}", text)))
    return max(markers, phone_count, 1)
